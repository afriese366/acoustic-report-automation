"""Liest das bewusst kleine und dokumentierte Demo-Excel-Schema."""

from pathlib import Path

from openpyxl import load_workbook

from acoustic_report.domain.models import (
    AcousticDataset,
    Contribution,
    InputValidationError,
    NoiseSource,
    Receiver,
)


SHEETS = {
    "Sources": ("source_id", "description", "sound_power_db", "correction_db"),
    "Receivers": ("receiver_id", "description", "area_type", "limit_db"),
    "Contributions": ("receiver_id", "source_id", "level_db"),
}


def _records(sheet, expected_headers: tuple[str, ...]) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise InputValidationError(f"Tabellenblatt {sheet.title} ist leer.")
    headers = tuple(str(value or "").strip() for value in rows[0])
    if headers != expected_headers:
        raise InputValidationError(
            f"Tabellenblatt {sheet.title} hat unerwartete Spalten: {headers}"
        )
    return [
        dict(zip(headers, row))
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]


def read_workbook(path: str | Path) -> AcousticDataset:
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Eingabedatei nicht gefunden: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing = sorted(set(SHEETS) - set(workbook.sheetnames))
        if missing:
            raise InputValidationError("Fehlende Tabellenblätter: " + ", ".join(missing))
        source_rows = _records(workbook["Sources"], SHEETS["Sources"])
        receiver_rows = _records(workbook["Receivers"], SHEETS["Receivers"])
        contribution_rows = _records(workbook["Contributions"], SHEETS["Contributions"])
    finally:
        workbook.close()

    try:
        dataset = AcousticDataset(
            sources=tuple(
                NoiseSource(
                    source_id=str(row["source_id"]).strip(),
                    description=str(row["description"]).strip(),
                    sound_power_db=float(row["sound_power_db"]),
                    correction_db=float(row["correction_db"] or 0),
                )
                for row in source_rows
            ),
            receivers=tuple(
                Receiver(
                    receiver_id=str(row["receiver_id"]).strip(),
                    description=str(row["description"]).strip(),
                    area_type=str(row["area_type"]).strip(),
                    limit_db=float(row["limit_db"]),
                )
                for row in receiver_rows
            ),
            contributions=tuple(
                Contribution(
                    receiver_id=str(row["receiver_id"]).strip(),
                    source_id=str(row["source_id"]).strip(),
                    level_db=float(row["level_db"]),
                )
                for row in contribution_rows
            ),
        )
    except (TypeError, ValueError) as exc:
        raise InputValidationError("Die Arbeitsmappe enthält ungültige Zahlen oder Texte.") from exc
    dataset.validate()
    return dataset
