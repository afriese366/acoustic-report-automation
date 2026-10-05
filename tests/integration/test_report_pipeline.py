import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from acoustic_report.application.generate_report import generate_report


ROOT = Path(__file__).resolve().parents[2]


class ReportPipelineTest(unittest.TestCase):
    def test_synthetic_project_creates_complete_report(self):
        with tempfile.TemporaryDirectory() as temporary:
            config_source = ROOT / "examples" / "synthetic_project" / "project.json"
            config_text = config_source.read_text(encoding="utf-8").replace(
                '"../../output/demo_report.docx"',
                '"' + (Path(temporary) / "demo_report.docx").as_posix() + '"',
            )
            config_path = Path(temporary) / "project.json"
            config_path.write_text(config_text, encoding="utf-8")
            # Absolute paths avoid coupling the temporary config location to repository layout.
            config_text = config_path.read_text(encoding="utf-8")
            config_text = config_text.replace(
                '"acoustic_data.xlsx"',
                '"' + (ROOT / "examples" / "synthetic_project" / "acoustic_data.xlsx").as_posix() + '"',
            ).replace(
                '"../../resources/demo_template.docx"',
                '"' + (ROOT / "resources" / "demo_template.docx").as_posix() + '"',
            ).replace(
                '"../../resources/demo_citations.json"',
                '"' + (ROOT / "resources" / "demo_citations.json").as_posix() + '"',
            )
            config_path.write_text(config_text, encoding="utf-8")

            output = generate_report(config_path)
            self.assertTrue(output.is_file())
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertNotIn("{{", text)
            self.assertNotIn("[[RESULTS_TABLE]]", text)
            self.assertIn("DEMO-2026-001", text)
            self.assertGreaterEqual(len(document.tables), 1)
            with ZipFile(output) as archive:
                core = archive.read("docProps/core.xml").decode("utf-8")
            self.assertNotIn("lastModifiedBy", core)


if __name__ == "__main__":
    unittest.main()
