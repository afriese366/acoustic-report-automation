"""Erzeugt alle binären Demo-Artefakte ausschließlich aus synthetischen Werten."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]


def _set_font(style, name: str, size: int, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def create_template() -> Path:
    output = ROOT / "resources" / "demo_template.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    _set_font(document.styles["Normal"], "Aptos", 10)
    document.styles["Normal"].paragraph_format.space_after = Pt(4)
    document.styles["Normal"].paragraph_format.line_spacing = 1.0
    _set_font(document.styles["Title"], "Aptos Display", 22, bold=True)
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    _set_font(document.styles["Subtitle"], "Aptos", 11)
    document.styles["Subtitle"].font.italic = False
    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        _set_font(document.styles[style_name], "Aptos Display", size, bold=True)

    title = document.add_paragraph(style="Title")
    title.add_run("Synthetic Acoustic Assessment")
    subtitle = document.add_paragraph("Automated report generated from validated demonstration data")
    subtitle.style = document.styles["Subtitle"]

    document.add_paragraph("Project {{ project.id }}")
    document.add_paragraph("{{ project.title }}")
    document.add_paragraph("Report date {{ project.report_date }}")
    document.add_paragraph("Prepared by {{ project.author }}")

    document.add_heading("Purpose and scope", level=1)
    document.add_paragraph(
        "This report demonstrates a deterministic document automation pipeline. "
        "It validates structured input, applies transparent acoustic rules and produces a reviewable Word document."
    )

    document.add_heading("Input data", level=1)
    document.add_paragraph(
        "The demonstration model contains {{ sources|length }} synthetic sound sources. "
        "All identifiers, descriptions and levels were created for this portfolio application."
    )
    document.add_paragraph("{%p for source in sources %}")
    document.add_paragraph(
        "{{ source.id }}  {{ source.description }}  sound power {{ source.sound_power }}, "
        "correction {{ source.correction }} dB, corrected value {{ source.corrected }}"
    )
    document.add_paragraph("{%p endfor %}")

    document.add_heading("Assessment method", level=1)
    document.add_paragraph(
        "Contributions at each receiver are combined energetically. The calculated level is rounded "
        "commercially to a whole decibel and compared with the assigned daytime limit. "
        "The method and validation contract are documented in {{ citations[0].reference }} and {{ citations[1].reference }}."
    )

    document.add_heading("Results", level=1)
    document.add_paragraph("[[RESULTS_TABLE]]")

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        "{{ summary.compliant_count }} of {{ summary.receiver_count }} assessed receivers meet their assigned limits."
    )
    document.add_paragraph("{%p if summary.all_compliant %}")
    document.add_paragraph(
        "The synthetic scenario meets all configured assessment criteria."
    )
    document.add_paragraph("{%p else %}")
    document.add_paragraph(
        "At least one synthetic receiver exceeds its configured criterion and requires review."
    )
    document.add_paragraph("{%p endif %}")

    document.add_heading("References", level=1)
    document.add_paragraph("{%p for citation in citations %}")
    document.add_paragraph("{{ citation.reference }} {{ citation.title }}")
    document.add_paragraph("{%p endfor %}")

    footer = section.footer.paragraphs[0]
    footer.text = "Acoustic Report Automation  Synthetic portfolio report"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "Synthetic Acoustic Assessment Template"
    properties.subject = "Portfolio demonstration"
    properties.keywords = ""
    properties.comments = ""
    document.save(output)
    return output


def create_workbook() -> Path:
    output = ROOT / "examples" / "synthetic_project" / "acoustic_data.xlsx"
    workbook = Workbook()
    sources = workbook.active
    sources.title = "Sources"
    sources.append(("source_id", "description", "sound_power_db", "correction_db"))
    sources.append(("S1", "Loading area", 92.4, 0.0))
    sources.append(("S2", "Rooftop ventilation", 84.8, 2.0))

    receivers = workbook.create_sheet("Receivers")
    receivers.append(("receiver_id", "description", "area_type", "limit_db"))
    receivers.append(("R1", "Birch Street 12", "General residential", 55.0))
    receivers.append(("R2", "Workshop Lane 4", "Mixed use", 60.0))

    contributions = workbook.create_sheet("Contributions")
    contributions.append(("receiver_id", "source_id", "level_db"))
    contributions.append(("R1", "S1", 49.2))
    contributions.append(("R1", "S2", 44.1))
    contributions.append(("R2", "S1", 52.5))
    contributions.append(("R2", "S2", 46.0))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 32)
    workbook.save(output)
    return output


if __name__ == "__main__":
    print(create_template())
    print(create_workbook())
